import sys
import tempfile
import zipfile
from pathlib import Path

import docx

src_path = str(Path(__file__).resolve().parent.parent / "src")
if src_path not in sys.path:
    sys.path.insert(0, src_path)

from craw.export_docx import export_markdown_to_docx


SAMPLE = {
    "topic": "制造业薪酬体系对标",
    "executive_summary": "对标样本显示关键岗位薪酬带宽需要重估。",
    "key_findings": [{"title": "带宽", "details": "操作岗与技师岗差距扩大。"}],
    "risk_assessment": ["同业数据口径不一致。"],
    "strategic_recommendations": ["先统一岗位序列再调薪。"],
    "sources": [
        {
            "title": "人社部工资价位",
            "url": "https://example.com/hr/salary-benchmark",
            "evidence": "给出分位值。",
        }
    ],
}


def _docx_text(path: Path) -> str:
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_structured_sources_become_links_in_md_and_docx():
    with tempfile.TemporaryDirectory() as tmp:
        target = Path(tmp) / "report.docx"
        export_markdown_to_docx(SAMPLE, target, auto_open=False)
        url = "https://example.com/hr/salary-benchmark"
        assert url in _docx_text(target)
        xml = zipfile.ZipFile(target).read("word/document.xml").decode("utf-8")
        assert "hyperlink" in xml
        assert "example.com/hr/salary-benchmark" in xml
        markdown = target.with_suffix(".md").read_text(encoding="utf-8")
        assert f"]({url})" in markdown
        assert "引用来源" in markdown or "来源" in markdown


def test_bare_urls_are_listed_as_sources():
    with tempfile.TemporaryDirectory() as tmp:
        source_md = Path(tmp) / "notes.md"
        source_md.write_text(
            "# 绩效考核体系从 KPI 向 OKR 过渡\n\n"
            "公开资料见 https://example.com/hr/okr-guide 以及 https://gov.example/labor/rules 。\n",
            encoding="utf-8",
        )
        target = Path(tmp) / "out.docx"
        export_markdown_to_docx(source_md, target, auto_open=False)
        rewritten = source_md.read_text(encoding="utf-8")
        assert "https://example.com/hr/okr-guide" in rewritten
        assert "https://gov.example/labor/rules" in rewritten
        assert "来源" in rewritten
        text = _docx_text(target)
        assert "https://example.com/hr/okr-guide" in text
        assert "https://gov.example/labor/rules" in text


if __name__ == "__main__":
    test_structured_sources_become_links_in_md_and_docx()
    test_bare_urls_are_listed_as_sources()
    print("引用来源导出测试通过")
