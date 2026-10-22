import sys
import tempfile
from pathlib import Path
from pptx import Presentation
import docx

src_path = str(Path(__file__).resolve().parent.parent / "src")
if src_path not in sys.path:
    sys.path.insert(0, src_path)

from craw.export_docx import get_report_paths, export_markdown_to_docx
from craw.export_pptx import export_report_to_pptx
from craw.export_charts import render_consulting_chart, extract_chart_data_from_report

SAMPLE_DATA = {
    "topic": "自动化测试课题：具身智能发展趋势",
    "executive_summary": "具身智能（Embodied AI）正在迎来从单体算法向物理实体软硬件协同的重要转折期。",
    "key_findings": [
        {"title": "硬件演进", "details": "端侧芯片能效比大幅提高，支持复杂物理交互推理。"},
        {"title": "数据闭环", "details": "合成数据与仿真平台（如 Omniverse）降低了训练采集成本。"},
        {"title": "商业落地", "details": "工业仓储与精密分拣成为最先规模化落地的标杆场景。"}
    ],
    "risk_assessment": [
        "算力瓶颈与长尾场景反应迟滞风险。",
        "复杂交互下的物理安全规避标准缺失。"
    ],
    "strategic_recommendations": [
        "建议构建标准化高保真物理仿真训练环境。",
        "推进软硬件协同设计与低功耗神经网络处理单元（NPU）适配。"
    ],
    "quantitative_chart": {
        "chart_title": "2024-2030年具身智能机器人市场渗透率预测",
        "chart_type": "line",
        "x_axis_label": "年份",
        "y_axis_label": "市场渗透率（%）",
        "data_points": [
            {"label": "2024", "value": 2.5, "unit": "%"},
            {"label": "2026E", "value": 8.0, "unit": "%"},
            {"label": "2028E", "value": 18.5, "unit": "%"},
            {"label": "2030E", "value": 36.0, "unit": "%"}
        ],
        "takeaway": "工业制造与复杂服务场景渗透加速，2030年预期突破36%渗透拐点。"
    },
    "sources": [
        {"title": "Nature Machine Intelligence 2026", "url": "https://nature.com/articles/embodied-ai", "evidence": "展示了仿真到真机的高泛化效率。"}
    ]
}


def test_get_report_paths():
    paths = get_report_paths("测试课题")
    assert len(paths) == 6
    archive_md, archive_docx, archive_pptx, latest_md, latest_docx, latest_pptx = paths
    assert archive_md.suffix == ".md"
    assert archive_docx.suffix == ".docx"
    assert archive_pptx.suffix == ".pptx"
    assert latest_md.name == "research_report.md"
    assert latest_docx.name == "research_report.docx"
    assert latest_pptx.name == "research_report.pptx"


def test_chart_rendering():
    with tempfile.TemporaryDirectory() as tmpdir:
        target_img = Path(tmpdir) / "chart.png"
        chart_data = extract_chart_data_from_report(SAMPLE_DATA)
        assert chart_data is not None
        render_consulting_chart(chart_data, target_img)
        assert target_img.exists()
        assert target_img.stat().st_size > 5000


def test_export_pptx_with_chart():
    with tempfile.TemporaryDirectory() as tmpdir:
        target_pptx = Path(tmpdir) / "test_report.pptx"
        exported = export_report_to_pptx(SAMPLE_DATA, target_pptx, auto_open=False)
        assert exported is not None
        assert target_pptx.exists()
        assert target_pptx.stat().st_size > 0

        # 验证幻灯片有效性与比例（封面 + 摘要 + 洞察 + 图表页 + 风险 + 建议 + 引用 = 7 页）
        prs = Presentation(str(target_pptx))
        assert len(prs.slides) >= 6
        # 16:9 比例校验
        assert round(prs.slide_width / prs.slide_height, 2) == 1.78


def test_export_docx_with_chart():
    with tempfile.TemporaryDirectory() as tmpdir:
        target_docx = Path(tmpdir) / "test_report.docx"
        exported = export_markdown_to_docx(SAMPLE_DATA, target_docx, auto_open=False)
        assert exported is not None
        assert target_docx.exists()
        assert target_docx.stat().st_size > 0

        doc = docx.Document(str(target_docx))
        assert len(doc.paragraphs) > 0
        # 验证是否生成了关联图表文件
        target_png = target_docx.with_suffix(".png")
        assert target_png.exists()


if __name__ == "__main__":
    test_get_report_paths()
    test_chart_rendering()
    test_export_pptx_with_chart()
    test_export_docx_with_chart()
    print("🎉 自动化导出与图表渲染测试全部通过！")
