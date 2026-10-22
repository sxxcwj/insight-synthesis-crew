import json
import re
import platform
import subprocess
from pathlib import Path
from typing import Dict, Any, List, Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from craw.export_charts import extract_chart_data_from_report, render_consulting_chart


# 商务配色板 (McKinsey / BCG 风格冷调深邃体系)
COLOR_BG_DARK = RGBColor(0x0F, 0x17, 0x2A)      # Slate 900
COLOR_BG_CARD = RGBColor(0xF8, 0xFA, 0xFC)      # Slate 50
COLOR_BORDER = RGBColor(0xE2, 0xE8, 0xF0)       # Slate 200
COLOR_PRIMARY = RGBColor(0x1D, 0x4E, 0xD8)      # Royal Blue
COLOR_TEXT_MAIN = RGBColor(0x0F, 0x17, 0x2A)    # Deep Slate
COLOR_TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)   # Muted Gray
COLOR_ALERT_RED = RGBColor(0xB9, 0x1C, 0x1C)    # Warning Red
COLOR_SUCCESS_GREEN = RGBColor(0x04, 0x78, 0x57)# Strategic Green
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_slide_header(slide, title: str, category: str = "INSIGHT SYNTHESIS CREW · 战略研究简报"):
    """为内容幻灯片添加标准商业级页头"""
    # 顶部微型分类标签
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.3))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.name = "Microsoft YaHei"
    p_cat.font.size = Pt(9.5)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_PRIMARY

    # 主大标题
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.5), Inches(0.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.name = "Microsoft YaHei"
    p_title.font.size = Pt(20)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_MAIN

    # 装饰细分割线
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_BORDER
    line.line.color.rgb = COLOR_BORDER


def create_cover_slide(prs, topic: str):
    """幻灯片 1：高级封面页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 顶部装饰彩条
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12)
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLOR_PRIMARY
    accent_bar.line.color.rgb = COLOR_PRIMARY

    # 机构 / 系统标签
    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.6), Inches(2.8), Inches(0.4)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xFF)
    badge.line.color.rgb = RGBColor(0xC7, 0xD2, 0xFE)
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "EXECUTIVE RESEARCH REPORT"
    p_b.alignment = PP_ALIGN.CENTER
    p_b.font.name = "Microsoft YaHei"
    p_b.font.size = Pt(9)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_PRIMARY

    # 课题大标题
    t_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.8), Inches(2.2))
    tf = t_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = topic
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_MAIN

    # 副标题与说明
    sub_box = slide.shapes.add_textbox(Inches(1.2), Inches(4.5), Inches(10.8), Inches(0.8))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "长伴咨询 · 依据检索底稿整理的研判，请核对来源页中的链接"
    p_sub.font.name = "Microsoft YaHei"
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED

    # 底部元数据卡片
    meta_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(5.8), Inches(10.9), Inches(0.7)
    )
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = COLOR_BG_CARD
    meta_box.line.color.rgb = COLOR_BORDER
    tf_m = meta_box.text_frame
    p_m = tf_m.paragraphs[0]
    p_m.text = "驱动引擎: DeepSeek-V4.1-Flash  |  框架体系: CrewAI Multi-Agent  |  交付规格: 咨询级战略洞察"
    p_m.font.name = "Microsoft YaHei"
    p_m.font.size = Pt(10)
    p_m.font.color.rgb = COLOR_TEXT_MUTED


def create_executive_summary_slide(prs, summary: str):
    """幻灯片 2：执行摘要 (Executive Summary)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "执行摘要与全景速览 (Executive Summary)")

    # 主核心提炼卡片容器
    card = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(4.8)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_BG_CARD
    card.line.color.rgb = COLOR_BORDER

    # 左侧亮蓝强调竖边条
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(0.12), Inches(4.8)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_PRIMARY
    accent.line.color.rgb = COLOR_PRIMARY

    # 卡片内部正文
    tb = slide.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p_badge = tf.paragraphs[0]
    p_badge.text = "📋 核心战略结论与研判底稿"
    p_badge.font.name = "Microsoft YaHei"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_PRIMARY

    p_body = tf.add_paragraph()
    p_body.space_before = Pt(14)
    p_body.line_spacing = 1.35
    p_body.text = summary
    p_body.font.name = "Microsoft YaHei"
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = COLOR_TEXT_MAIN


def create_key_findings_slide(prs, findings: List[Dict[str, Any]]):
    """幻灯片 3：核心洞察与发现 (3 列卡片布局)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "一、 核心洞察与多源核验发现 (Key Insights)")

    col_width = Inches(3.75)
    gap = Inches(0.24)
    left_margin = Inches(0.8)
    top_pos = Inches(1.6)
    height = Inches(5.2)

    # 限制最多渲染 3 列
    items = findings[:3]
    for idx, item in enumerate(items):
        cur_left = left_margin + idx * (col_width + gap)

        # 背景白卡
        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, cur_left, top_pos, col_width, height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = COLOR_BORDER

        # 顶部蓝色标号条
        header_tag = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, cur_left, top_pos, col_width, Inches(0.4)
        )
        header_tag.fill.solid()
        header_tag.fill.fore_color.rgb = COLOR_PRIMARY
        header_tag.line.color.rgb = COLOR_PRIMARY
        tf_tag = header_tag.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = f"INSIGHT 0{idx+1}"
        p_tag.alignment = PP_ALIGN.CENTER
        p_tag.font.name = "Microsoft YaHei"
        p_tag.font.size = Pt(9.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_WHITE

        # 标题与内容文本
        tb = slide.shapes.add_textbox(cur_left + Inches(0.2), top_pos + Inches(0.55), col_width - Inches(0.4), height - Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = item.get("title", f"洞察 {idx+1}")
        p_t.font.name = "Microsoft YaHei"
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_MAIN

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(8)
        p_d.line_spacing = 1.2
        # 截断避免文本溢出，保留核心精要
        details = item.get("details", "")
        p_d.text = (details[:260] + "...") if len(details) > 260 else details
        p_d.font.name = "Microsoft YaHei"
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = RGBColor(0x33, 0x41, 0x55)


def create_risk_assessment_slide(prs, risks: List[str]):
    """幻灯片 4：潜在风险与不确定性研判 (警告红调卡片)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "二、 潜在风险与不确定性研判 (Risk Assessment)")

    top_start = Inches(1.6)
    row_height = Inches(1.2)
    row_gap = Inches(0.18)

    # 渲染前 3-4 个关键风险
    for idx, risk_text in enumerate(risks[:4]):
        cur_top = top_start + idx * (row_height + row_gap)

        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), cur_top, Inches(11.733), row_height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0xFF, 0xFA, 0xFA)  # 淡淡的粉白底
        card.line.color.rgb = RGBColor(0xFE, 0xCD, 0xCD)      # 柔和红边框

        # 左侧红条
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), cur_top, Inches(0.1), row_height
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = COLOR_ALERT_RED
        bar.line.color.rgb = COLOR_ALERT_RED

        tb = slide.shapes.add_textbox(Inches(1.1), cur_top + Inches(0.1), Inches(11.2), row_height - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.line_spacing = 1.2
        r_tag = p.add_run()
        r_tag.text = f"⚠️ 风险点 0{idx+1}： "
        r_tag.font.name = "Microsoft YaHei"
        r_tag.font.bold = True
        r_tag.font.size = Pt(10.5)
        r_tag.font.color.rgb = COLOR_ALERT_RED

        r_text = p.add_run()
        r_text.text = (risk_text[:200] + "...") if len(risk_text) > 200 else risk_text
        r_text.font.name = "Microsoft YaHei"
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = COLOR_TEXT_MAIN


def create_recommendations_slide(prs, recs: List[str]):
    """幻灯片 5：前瞻性战略行动建议 (决策绿调卡片)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "三、 前瞻性战略行动建议 (Strategic Recommendations)")

    top_start = Inches(1.6)
    row_height = Inches(1.2)
    row_gap = Inches(0.18)

    # 渲染前 3-4 项行动建议
    for idx, rec_text in enumerate(recs[:4]):
        cur_top = top_start + idx * (row_height + row_gap)

        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), cur_top, Inches(11.733), row_height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0xFA, 0xFD, 0xFA)  # 淡淡的浅绿底
        card.line.color.rgb = RGBColor(0xBB, 0xF7, 0xD0)      # 浅绿边框

        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), cur_top, Inches(0.1), row_height
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = COLOR_SUCCESS_GREEN
        bar.line.color.rgb = COLOR_SUCCESS_GREEN

        tb = slide.shapes.add_textbox(Inches(1.1), cur_top + Inches(0.1), Inches(11.2), row_height - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.line_spacing = 1.2
        r_tag = p.add_run()
        r_tag.text = f"ACTION 0{idx+1} 决策路线： "
        r_tag.font.name = "Microsoft YaHei"
        r_tag.font.bold = True
        r_tag.font.size = Pt(10.5)
        r_tag.font.color.rgb = COLOR_SUCCESS_GREEN

        r_text = p.add_run()
        r_text.text = (rec_text[:200] + "...") if len(rec_text) > 200 else rec_text
        r_text.font.name = "Microsoft YaHei"
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = COLOR_TEXT_MAIN


def create_sources_slide(prs, sources: List[Dict[str, Any]]):
    """幻灯片 6：权威引用与数据源清单"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "四、 权威文献与一手数据出处 (References & Sources)")

    top_start = Inches(1.6)
    row_height = Inches(0.85)
    row_gap = Inches(0.14)

    for idx, src in enumerate(sources[:5]):
        cur_top = top_start + idx * (row_height + row_gap)

        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), cur_top, Inches(11.733), row_height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = COLOR_BORDER

        tb = slide.shapes.add_textbox(Inches(1.0), cur_top + Inches(0.08), Inches(11.3), row_height - Inches(0.16))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        title = src.get("title", f"信源 {idx+1}") if isinstance(src, dict) else str(src)
        url = src.get("url", "") if isinstance(src, dict) else ""
        evidence = src.get("evidence", "") if isinstance(src, dict) else ""

        r_idx = p.add_run()
        r_idx.text = f"[{idx+1}] {title}  "
        r_idx.font.name = "Microsoft YaHei"
        r_idx.font.bold = True
        r_idx.font.size = Pt(9.5)
        r_idx.font.color.rgb = COLOR_TEXT_MAIN

        if url:
            r_url = p.add_run()
            r_url.text = f"({url[:65]}...)\n" if len(url) > 65 else f"({url})\n"
            r_url.font.name = "Calibri"
            r_url.font.size = Pt(8.5)
            r_url.font.color.rgb = COLOR_PRIMARY
            r_url.font.underline = True

        if evidence:
            r_ev = p.add_run()
            r_ev.text = f"事实支撑: {evidence[:130]}..."
            r_ev.font.name = "Microsoft YaHei"
            r_ev.font.size = Pt(8.5)
            r_ev.font.color.rgb = COLOR_TEXT_MUTED


def create_quantitative_chart_slide(prs, chart_info: Dict[str, Any], chart_image_path: Path):
    """幻灯片：核心量化指标与发展趋势研判 (图文对照)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = chart_info.get("chart_title", "核心量化发展趋势研判")
    add_slide_header(slide, title, category="QUANTITATIVE FORECAST & TRENDS · 量化趋势")

    # 左侧嵌入高分辨率图表
    slide.shapes.add_picture(str(chart_image_path), Inches(0.8), Inches(1.55), Inches(7.6), Inches(4.7))

    # 右侧卡片：核心数据亮点与战略启示
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.7), Inches(1.55), Inches(3.8), Inches(4.7)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_BG_CARD
    card.line.color.rgb = COLOR_BORDER
    card.line.width = Pt(1)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.3)
    tf.margin_bottom = Inches(0.3)
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)

    p0 = tf.paragraphs[0]
    p0.text = "📊 数据透视与战略启示"
    p0.font.name = "Microsoft YaHei"
    p0.font.size = Pt(12.5)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_PRIMARY
    p0.space_after = Pt(12)

    takeaway = chart_info.get("takeaway", "")
    if takeaway:
        p_tk = tf.add_paragraph()
        r_tk_title = p_tk.add_run()
        r_tk_title.text = "核心研判结论：\n"
        r_tk_title.font.bold = True
        r_tk_title.font.size = Pt(10)
        r_tk_title.font.color.rgb = COLOR_TEXT_MAIN
        r_tk = p_tk.add_run()
        r_tk.text = f"{takeaway}\n"
        r_tk.font.size = Pt(9.5)
        r_tk.font.color.rgb = COLOR_TEXT_MUTED
        p_tk.space_after = Pt(10)

    points = chart_info.get("data_points", [])
    if points:
        p_pts = tf.add_paragraph()
        r_pts_title = p_pts.add_run()
        r_pts_title.text = "关键节点观测值：\n"
        r_pts_title.font.bold = True
        r_pts_title.font.size = Pt(10)
        r_pts_title.font.color.rgb = COLOR_TEXT_MAIN
        for pt in points[:5]:
            p_item = tf.add_paragraph()
            p_item.text = f"• {pt.get('label', '')}: {pt.get('value', '')}{pt.get('unit', '')}"
            p_item.font.size = Pt(9)
            p_item.font.color.rgb = COLOR_TEXT_MAIN
            p_item.space_after = Pt(3)


def export_report_to_pptx(md_or_data: Any, pptx_path: Path, chart_image_path: Optional[Path] = None, auto_open: bool = True):
    """
    将研究报告解析并导出为 16:9 咨询级宽屏 PowerPoint (.pptx)
    并自动调用系统 Microsoft PowerPoint 打开
    """
    data = None
    if isinstance(md_or_data, dict):
        data = md_or_data
    elif isinstance(md_or_data, Path) or isinstance(md_or_data, str):
        content = Path(md_or_data).read_text(encoding="utf-8")
        # 优先解析 JSON
        m = re.search(r"\{[\s\S]*\}", content)
        if m:
            try:
                data = json.loads(m.group(0))
            except Exception:
                data = None

        # 正则从 Markdown 解析
        if not data:
            topic_m = re.search(r"^#\s+(.+)$", content, re.MULTILINE)
            topic_val = topic_m.group(1).strip() if topic_m else "研究课题简报"

            exec_m = re.search(r">\s*\*\*📋[^\*]+\*\*\s*\n+>\s*(.+?)(?=\n##|\Z)", content, re.DOTALL)
            exec_val = exec_m.group(1).strip().replace("\n> ", " ").replace("\n>", "") if exec_m else ""

            findings = []
            f_matches = re.findall(r"###\s*1\.\d+\s*([^\n]+)\n+([\s\S]+?)(?=\n###|\n##|\Z)", content)
            for ft, fd in f_matches:
                findings.append({"title": ft.strip(), "details": fd.strip()})

            risks = re.findall(r"\*\s*\*\*⚠️[^\*]+\*\*：?\s*([^\n]+)", content)

            recs = []
            rec_sec = re.search(r"##\s*三[^\n]+\n+([\s\S]+?)(?=\n##|\Z)", content)
            if rec_sec:
                recs = re.findall(r"^\d+\.\s*([^\n]+)", rec_sec.group(1), re.MULTILINE)

            sources = []
            src_sec = re.search(r"##\s*四[^\n]+\n+([\s\S]+?)(?=\n##|\Z)", content)
            if src_sec:
                for s in re.findall(r"^\d+\.\s*\[([^\]]+)\]\(([^)]+)\)\s*(?:——|-)?\s*([^\n]*)", src_sec.group(1), re.MULTILINE):
                    sources.append({"title": s[0].strip(), "url": s[1].strip(), "evidence": s[2].strip()})

            data = {
                "topic": topic_val,
                "executive_summary": exec_val,
                "key_findings": findings,
                "risk_assessment": risks,
                "strategic_recommendations": recs,
                "sources": sources
            }

    if not data:
        print("未提取到有效报告数据，跳过 PPT 导出")
        return None

    # 尝试自动提取量化图表数据并生成高清图表
    chart_info = None
    if not chart_image_path:
        chart_info = extract_chart_data_from_report(data or md_or_data)
        if chart_info:
            c_path = pptx_path.with_suffix(".png")
            try:
                render_consulting_chart(chart_info, c_path)
                chart_image_path = c_path
            except Exception as e:
                print(f"图表渲染提示: {e}")
    elif chart_image_path and Path(chart_image_path).exists():
        chart_info = extract_chart_data_from_report(data or md_or_data) or {
            "chart_title": "核心量化发展趋势研判",
            "takeaway": "量化数据印证了该领域的强劲演进趋势与核心支撑力。",
            "data_points": []
        }

    # 初始化 16:9 演示文稿
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    topic = data.get("topic", "战略研究洞察报告")
    summary = data.get("executive_summary", "")
    findings = data.get("key_findings", [])
    risks = data.get("risk_assessment", [])
    recs = data.get("strategic_recommendations", [])
    sources = data.get("sources", [])

    # 生成各页幻灯片
    create_cover_slide(prs, topic)
    if summary:
        create_executive_summary_slide(prs, summary)
    if findings:
        create_key_findings_slide(prs, findings)
    if chart_image_path and Path(chart_image_path).exists() and chart_info:
        create_quantitative_chart_slide(prs, chart_info, chart_image_path)
    if risks:
        create_risk_assessment_slide(prs, risks)
    if recs:
        create_recommendations_slide(prs, recs)
    if sources:
        create_sources_slide(prs, sources)

    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(pptx_path))
    print(f"🖥️ [演示文稿] 咨询级 16:9 汇报 PPT 已生成: {pptx_path}")

    # 遵循用户系统级规则：为用户生成 PPT 后，自动调用系统命令使用 Microsoft PowerPoint 打开
    if auto_open and platform.system() == "Darwin":
        try:
            res = subprocess.run(["open", "-a", "Microsoft PowerPoint", str(pptx_path)], check=False, capture_output=True)
            if res.returncode != 0:
                subprocess.run(["open", str(pptx_path)], check=False)
            print("🚀 已自动调用 Microsoft PowerPoint 为您直接打开汇报幻灯片！")
        except Exception as e:
            print(f"自动打开 PPT 提示: {e}")

    return pptx_path


if __name__ == "__main__":
    report_md = Path("reports/research_report.md")
    report_pptx = Path("reports/research_report.pptx")
    if report_md.exists():
        export_report_to_pptx(report_md, report_pptx)
